"""Split combined demo notes into bank-only and credit-union-only .docx copies."""
from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "tools" / "_demo_master.docx"
OUT_BANK = ROOT / "L1_Executive_Dashboard_Suite_Demo_Notes_Bank.docx"
OUT_CU = ROOT / "L1_Executive_Dashboard_Suite_Demo_Notes_Credit_Union.docx"

CU_TARGETS = [
    ("Level 1-2", "Return on Assets", "0.75%"),
    ("Level 1-2", "Net worth growth", "8.0%"),
    ("Level 1-2", "Net Interest Margin", "3.50%"),
    ("Level 1-2", "Efficiency ratio", "79% (lower is better, vs Financial Performance Report)"),
    ("Level 1-2", "Net Operating Income versus budget", "$12.5 million"),
    ("Level 1-2", "Loan growth", "6.0% year-over-year"),
    ("Level 1-2", "Share growth", "5.5% year-over-year"),
    ("Level 1-2", "Noninterest income mix", "16%"),
    ("Level 1-3", "Net Promoter Score", "70"),
    ("Level 1-3", "Retention", "93%"),
    ("Level 1-3", "Member growth", "3.5%"),
    ("Level 1-3", "Core share mix", "80%"),
    ("Level 1-3", "Average Member Relationship", "$18.5k"),
    ("Level 1-3", "Digital active", "58%"),
    ("Level 1-3", "Cross-sell", "16%"),
    ("Level 1-3", "Customer Acquisition Cost", "$240"),
    ("Level 1-4", "Net worth ratio", "10.5% floor"),
    ("Level 1-4", "Past due 30 days or more", "0.90% limit"),
    ("Level 1-4", "Criticized assets", "3.5% limit"),
    ("Level 1-4", "Allowance for Credit Losses coverage", "1.20×"),
    ("Level 1-4", "Loans-to-Shares", "82% plan; band 75–88%"),
    ("Level 1-4", "Interest Rate Risk net interest income +200 basis points", "−5.5% plan; −8% floor"),
    ("Level 1-4", "Wholesale reliance", "10% limit"),
    ("Level 1-4", "Bank Secrecy Act aging", "16 days average"),
    ("Level 1-5", "Decision cycle", "12 days"),
    ("Level 1-5", "Funnel conversion", "42%"),
    ("Level 1-5", "Yield on new originations", "6.40%"),
    ("Level 1-5", "Origination volume", "$22 million / month"),
    ("Level 1-5", "Cost per transaction", "$2.10"),
    ("Level 1-5", "Digital self-service", "68%"),
    ("Level 1-5", "Onboarding cycle", "1.2 days"),
    ("Level 1-5", "Teller service-level agreement", "1.8 minutes"),
    ("Level 1-6", "Engagement", "76 (Engaged band)"),
    ("Level 1-6", "Voluntary turnover", "12%"),
    ("Level 1-6", "Teller / frontline turnover", "18%"),
    ("Level 1-6", "Time-to-fill critical roles", "50 days"),
    ("Level 1-6", "Required training completion", "98%"),
    ("Level 1-6", "Succession coverage", "75%"),
    ("Level 1-6", "Board education hours", "12 hours"),
]

SLASH_BANK = {
    "Return on Equity  /  Net worth growth": "Return on Equity",
    "Return on Equity / Net worth growth": "Return on Equity",
    "Customer Satisfaction Index / Net Promoter Score": "Customer Satisfaction Index",
    "Customer Satisfaction Index / Net Promoter Score  /  Net Promoter Score": "Customer Satisfaction Index",
    "Deposit growth  /  Share growth": "Deposit growth",
    "Deposit growth / Share growth": "Deposit growth",
    "Customer growth  /  Member growth": "Customer growth",
    "Customer growth / Member growth": "Customer growth",
    "Core funding mix  /  Core share mix": "Core funding mix",
    "Core funding mix / Core share mix": "Core funding mix",
    "Products per household  /  Average Member Relationship": "Products per household",
    "Products per household / Average Member Relationship": "Products per household",
    "Tier 1 leverage  /  Net worth ratio": "Tier 1 leverage",
    "Tier 1 leverage / Net worth ratio": "Tier 1 leverage",
    "Loans-to-Deposits  /  Loans-to-Shares": "Loans-to-Deposits",
    "Loans-to-Deposits / Loans-to-Shares": "Loans-to-Deposits",
    "Commercial pipeline  /  Onboarding cycle": "Commercial pipeline",
    "Commercial pipeline / Onboarding cycle": "Commercial pipeline",
    "Commercial skill coverage  /  Board education hours": "Commercial skill coverage",
    "Commercial skill coverage / Board education hours": "Commercial skill coverage",
    "Customer / Member Franchise": "Customer Franchise",
    "Customer / member franchise": "Customer franchise",
    "Metric (bank / credit union)": "Metric",
}

SLASH_CU = {
    "Return on Equity  /  Net worth growth": "Net worth growth",
    "Return on Equity / Net worth growth": "Net worth growth",
    "Customer Satisfaction Index / Net Promoter Score": "Net Promoter Score",
    "Customer Satisfaction Index / Net Promoter Score  /  Net Promoter Score": "Net Promoter Score",
    "Deposit growth  /  Share growth": "Share growth",
    "Deposit growth / Share growth": "Share growth",
    "Customer growth  /  Member growth": "Member growth",
    "Customer growth / Member growth": "Member growth",
    "Core funding mix  /  Core share mix": "Core share mix",
    "Core funding mix / Core share mix": "Core share mix",
    "Products per household  /  Average Member Relationship": "Average Member Relationship",
    "Products per household / Average Member Relationship": "Average Member Relationship",
    "Tier 1 leverage  /  Net worth ratio": "Net worth ratio",
    "Tier 1 leverage / Net worth ratio": "Net worth ratio",
    "Loans-to-Deposits  /  Loans-to-Shares": "Loans-to-Shares",
    "Loans-to-Deposits / Loans-to-Shares": "Loans-to-Shares",
    "Commercial pipeline  /  Onboarding cycle": "Onboarding cycle",
    "Commercial pipeline / Onboarding cycle": "Onboarding cycle",
    "Commercial skill coverage  /  Board education hours": "Board education hours",
    "Commercial skill coverage / Board education hours": "Board education hours",
    "Customer / Member Franchise": "Member Franchise",
    "Customer / member franchise": "Member franchise",
    "Metric (bank / credit union)": "Metric",
}


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def pick_bank_cu_segments(text: str, charter: str) -> str:
    if "Credit union:" not in text and "Bank:" not in text:
        return text
    if charter == "bank":
        if re.search(r"\s*\|\s*Credit union:", text):
            text = re.split(r"\s*\|\s*Credit union:", text, maxsplit=1)[0]
        elif re.search(r"\nCredit union:", text):
            text = re.split(r"\nCredit union:", text, maxsplit=1)[0]
        text = re.sub(r"^Bank:\s*", "", text.strip())
        return text.strip()
    if re.search(r"Credit union:\s*", text):
        m = re.search(r"Credit union:\s*(.*)$", text, re.DOTALL)
        return m.group(1).strip() if m else text
    if text.strip().startswith("Bank:"):
        m = re.search(r"Bank:\s*(.*?)(?:\s*\|\s*Credit union:|\nCredit union:|$)", text, re.DOTALL)
        return m.group(1).strip() if m else text
    return text


def apply_slash_labels(text: str, charter: str) -> str:
    mapping = SLASH_BANK if charter == "bank" else SLASH_CU
    for combo, label in mapping.items():
        text = text.replace(combo, label)
    return text


def transform_text(text: str, charter: str) -> str:
    if not text or not text.strip():
        return text

    t = text

    if charter == "bank":
        t = t.replace("community bank or credit union", "community bank")
        t = t.replace("Customer / member franchise", "Customer franchise")
        t = t.replace("customer or members", "customers")
        t = t.replace("customers or members", "customers")
        t = t.replace("Customer / Member", "Customer")
        t = t.replace("household / member growth", "household growth")
        t = t.replace("products per household / Average Member Relationship", "products per household")
        t = t.replace("deposit / share growth", "deposit growth")
        t = t.replace("pipeline / onboarding", "commercial pipeline")
        t = t.replace("skill / board education", "commercial skill coverage")
        t = re.sub(
            r"Bank story:\s*(.*?)\.\s*Credit-union story:.*",
            r"\1.",
            t,
            flags=re.DOTALL,
        )
        t = re.sub(
            r"Banks measure household franchise depth \([^)]+\)\.\s*Credit unions measure member-owner depth \([^)]+\)\.\s*The math is the same spine; the labels change\.\s*",
            r"Banks measure household franchise depth (Customer Satisfaction Index, household retention, products per household, core deposit mix). ",
            t,
        )
        t = re.sub(
            r"Bank caption: People → Process → Customer → Financial \(shareholder value\)\.\s*Risk constrains\.\s*Credit-union caption: People → Process → Member → Financial \(member value \+ net worth\)\.\s*Risk constrains\.\s*Sample causal lines in the mockup: bank, credit quality → provision → Return on Assets;\s*credit union, efficiency → capital retention → net worth\.",
            "People → Process → Customer → Financial (shareholder value). Risk constrains. Sample causal line in the mockup: credit quality → provision → Return on Assets.",
            t,
        )
        t = re.sub(
            r"Sample causal lines in the mockup: bank, credit quality → provision → Return on Assets;\s*credit union, efficiency → capital retention → net worth\.",
            "Sample causal line in the mockup: credit quality → provision → Return on Assets.",
            t,
        )
        t = t.replace(
            "Capital generation (Return on Equity or net worth growth) is scored here. Capital adequacy (Tier 1 leverage or net worth ratio) lives on Level 1-4 and is echoed on Level 1-1.",
            "Return on Equity is scored here. Capital adequacy (Tier 1 leverage) lives on Level 1-4 and is echoed on Level 1-1.",
        )
        t = t.replace("products / Average Member Relationship (Level 1-3)", "products per household (Level 1-3)")
        t = t.replace("Loans-to-Deposits or Loans-to-Shares", "Loans-to-Deposits")
        t = t.replace("Loans-to-Deposits / Loans-to-Shares", "Loans-to-Deposits")
        t = re.sub(
            r"Operations \(Chief Operating Officer\): decision cycle, cost per transaction, digital self-service, teller service-level agreement, and — in credit-union mode — member onboarding cycle\. Production \(Chief Business Officer\): funnel conversion, yield on new originations, origination volume, commercial pipeline \(bank\)\.",
            "Operations (Chief Operating Officer): decision cycle, cost per transaction, digital self-service, teller service-level agreement. Production (Chief Business Officer): funnel conversion, yield on new originations, origination volume, commercial pipeline.",
            t,
        )
        t = t.replace(
            "Chief Operating Officer (operations: cycle time, cost per transaction, self-service, teller service-level agreement, credit-union onboarding)",
            "Chief Operating Officer (operations: cycle time, cost per transaction, self-service, teller service-level agreement)",
        )
        t = t.replace("Never mix Financial Performance Report curves on the same tile.", "Peer curves use Uniform Bank Performance Report only.")
        t = t.replace("Never mix Uniform Bank Performance Report and Financial Performance Report.", "Peer curves use Uniform Bank Performance Report only.")
        t = re.sub(
            r"Credit-union mode swaps commercial skill coverage for volunteer board education hours — a Management / governance metric, not a bank credential count\.\s*",
            "",
            t,
        )
        t = t.replace(
            "Peer chip: ahead of or behind the same-charter, same-size median. Uniform Bank Performance Report $1 billion–$3 billion for the bank; Financial Performance Report for the credit-union demo.",
            "Peer chip: ahead of or behind the Uniform Bank Performance Report median ($1 billion–$3 billion assets).",
        )
        t = t.replace(
            "Curves never mix Uniform Bank Performance Report and Financial Performance Report.",
            "Peer curves use Uniform Bank Performance Report only.",
        )
        t = t.replace(
            "Changes apply to the active charter side only (Bank versus Credit union), so Return on Equity and net worth growth stay independent.",
            "Changes apply to the bank charter sample in this session.",
        )
        t = t.replace(
            "In a live bank, earnings and franchise targets changed mid-year still go through Asset-Liability Committee.",
            "In a live institution, earnings and franchise targets changed mid-year still go through Asset-Liability Committee.",
        )
    else:
        t = t.replace("community bank or credit union", "credit union")
        t = t.replace("Riverbend Community Bank", "Riverbend Community Credit Union")
        t = t.replace("Customer / member franchise", "Member franchise")
        t = t.replace("Customer / Member", "Member")
        t = t.replace("customer or members", "members")
        t = t.replace("customers or members", "members")
        t = t.replace("household / member growth", "member growth")
        t = t.replace("products per household / Average Member Relationship", "Average Member Relationship")
        t = t.replace("deposit / share growth", "share growth")
        t = t.replace("pipeline / onboarding", "onboarding cycle")
        t = t.replace("skill / board education", "board education hours")
        t = re.sub(
            r"Bank story:\s*.*?\.\s*Credit-union story:\s*(.*?)\.\s*Capital generation \(Return on Equity or net worth growth\)",
            r"\1. Capital generation (net worth growth)",
            t,
            flags=re.DOTALL,
        )
        t = t.replace(
            "Capital generation (Return on Equity or net worth growth) is scored here. Capital adequacy (Tier 1 leverage or net worth ratio) lives on Level 1-4 and is echoed on Level 1-1.",
            "Net worth growth is scored here. Capital adequacy (net worth ratio) lives on Level 1-4 and is echoed on Level 1-1.",
        )
        t = re.sub(
            r"Bank caption: People → Process → Customer → Financial \(shareholder value\)\.\s*Risk constrains\.\s*Credit-union caption: People → Process → Member → Financial \(member value \+ net worth\)\.\s*Risk constrains\.\s*Sample causal lines in the mockup: bank, credit quality → provision → Return on Assets;\s*credit union, efficiency → capital retention → net worth\.",
            "People → Process → Member → Financial (member value + net worth). Risk constrains. Sample causal line in the mockup: efficiency → capital retention → net worth.",
            t,
        )
        t = re.sub(
            r"Banks measure household franchise depth \([^)]+\)\.\s*Credit unions measure member-owner depth \([^)]+\)\.\s*The math is the same spine; the labels change\.\s*",
            "Credit unions measure member-owner depth (Net Promoter Score, member retention, Average Member Relationship, core share mix). ",
            t,
        )
        t = t.replace("products / Average Member Relationship (Level 1-3)", "Average Member Relationship (Level 1-3)")
        t = t.replace("Loans-to-Deposits or Loans-to-Shares", "Loans-to-Shares")
        t = t.replace("Loans-to-Deposits / Loans-to-Shares", "Loans-to-Shares")
        t = re.sub(
            r"Operations \(Chief Operating Officer\): decision cycle, cost per transaction, digital self-service, teller service-level agreement, and — in credit-union mode — member onboarding cycle\. Production \(Chief Business Officer\): funnel conversion, yield on new originations, origination volume, commercial pipeline \(bank\)\.",
            "Operations (Chief Operating Officer): decision cycle, cost per transaction, digital self-service, teller service-level agreement, member onboarding cycle. Production (Chief Business Officer): funnel conversion, yield on new originations, origination volume.",
            t,
        )
        t = re.sub(
            r"Credit-union mode swaps commercial skill coverage for volunteer board education hours — a Management / governance metric, not a bank credential count\.",
            "Commercial skill coverage swaps to volunteer board education hours — a Management / governance metric.",
            t,
        )
        t = t.replace(
            "Peer chip: ahead of or behind the same-charter, same-size median. Uniform Bank Performance Report $1 billion–$3 billion for the bank; Financial Performance Report for the credit-union demo.",
            "Peer chip: ahead of or behind the Financial Performance Report median (5300 Call Report peer set).",
        )
        t = t.replace(
            "Curves never mix Uniform Bank Performance Report and Financial Performance Report.",
            "Peer curves use Financial Performance Report only.",
        )
        t = t.replace(
            "Changes apply to the active charter side only (Bank versus Credit union), so Return on Equity and net worth growth stay independent.",
            "Changes apply to the credit union charter sample in this session.",
        )
        t = t.replace(
            "Green / amber / red is not a Uniform Bank Performance Report grade.",
            "Green / amber / red is not a Financial Performance Report grade.",
        )
        t = t.replace(
            "Chief Operating Officer (operations: cycle time, cost per transaction, self-service, teller service-level agreement, credit-union onboarding)",
            "Chief Operating Officer (operations: cycle time, cost per transaction, self-service, teller service-level agreement, member onboarding cycle)",
        )
        t = t.replace("Never mix Uniform Bank Performance Report curves on the same tile.", "Peer curves use Financial Performance Report only.")
        t = t.replace("Never mix Uniform Bank Performance Report and Financial Performance Report.", "Peer curves use Financial Performance Report only.")

    t = apply_slash_labels(t, charter)
    t = pick_bank_cu_segments(t, charter)

    if charter == "bank":
        t = re.sub(r"\s*\(credit union\)[^.|]*", "", t, flags=re.IGNORECASE)
        t = t.replace("net worth ratio (credit union)", "Tier 1 leverage")
    else:
        t = re.sub(r"\s*\(bank\)[^.|]*", "", t, flags=re.IGNORECASE)

    return t


def transform_document(doc: Document, charter: str) -> None:
    to_delete: list[Paragraph] = []
    for paragraph in doc.paragraphs:
        raw = paragraph.text
        if charter == "bank" and raw.strip() == "Bank versus credit union":
            set_paragraph_text(paragraph, "Bank charter and peers")
            continue
        if charter == "cu" and raw.strip() == "Bank versus credit union":
            set_paragraph_text(paragraph, "Credit union charter and peers")
            continue
        if "One framework, two label sets. The Bank / Credit union toggle" in raw:
            if charter == "bank":
                set_paragraph_text(
                    paragraph,
                    "Open the six Level 1 HTML files with Bank selected (default). Peer context is Uniform Bank Performance Report (Call Report), assets $1 billion–$3 billion. Riverbend Community Bank is the sample institution. Do not mix Financial Performance Report curves on the same tile.",
                )
            else:
                set_paragraph_text(
                    paragraph,
                    "Select Credit union on the top bar before the walk. Peer context is Financial Performance Report / 5300 Call Report for the same asset band. Riverbend Community Credit Union is the demo label set. Do not mix Uniform Bank Performance Report curves on the same tile.",
                )
            continue
        if raw.startswith("Top bar: institution name, Bank / Credit union toggle"):
            if charter == "bank":
                set_paragraph_text(
                    paragraph,
                    "Top bar: institution name, Bank charter (default), four page pills (Monitor, Add, Edit, Definitions).",
                )
            else:
                set_paragraph_text(
                    paragraph,
                    "Top bar: institution name, Credit union charter (toggle from Bank), four page pills (Monitor, Add, Edit, Definitions).",
                )
            continue
        if raw.strip() == "Stay on Bank unless someone asks for the credit-union toggle.":
            if charter == "bank":
                to_delete.append(paragraph)
            else:
                set_paragraph_text(
                    paragraph,
                    "Select Credit union on the top bar before starting the walk.",
                )
            continue
        if "If asked about credit unions: toggle once on Level 1-2" in raw:
            if charter == "bank":
                to_delete.append(paragraph)
            else:
                set_paragraph_text(
                    paragraph,
                    "On Level 1-2, point out net worth growth (not Return on Equity) and efficiency versus the Financial Performance Report curve.",
                )
            continue
        if raw.startswith("12. Sample targets in the demo bank pack"):
            if charter == "cu":
                set_paragraph_text(paragraph, "12. Sample targets in the demo credit union pack")
            continue
        if "Credit-union demo targets differ" in raw:
            if charter == "bank":
                set_paragraph_text(
                    paragraph,
                    "These are workshop stand-ins so the mockups have bindable objectives. They are not Riverbend’s published policy.",
                )
            else:
                set_paragraph_text(
                    paragraph,
                    "These are workshop stand-ins for the credit union demo pack. They are not Riverbend’s published policy.",
                )
            continue
        if "Bank and credit-union wording both appear because the charter toggle" in raw:
            if charter == "bank":
                set_paragraph_text(
                    paragraph,
                    "Compiled from the Definitions tab in the Level 1 suite (bank charter). One row per proving metric. CAMELS letters are tags, not weights — the score is still Actual versus the institution target. Peer context is Uniform Bank Performance Report.",
                )
            else:
                set_paragraph_text(
                    paragraph,
                    "Compiled from the Definitions tab in the Level 1 suite (credit union charter). One row per proving metric. CAMELS letters are tags, not weights — the score is still Actual versus the institution target. Peer context is Financial Performance Report / 5300 Call Report.",
                )
            continue
        if "The Metrics column lists every proving tile" in raw and "Bank and credit-union names" in raw:
            if charter == "bank":
                set_paragraph_text(
                    paragraph,
                    "The Metrics column lists every proving tile on the six Level 1 pages (from the dashboard files), with abbreviations spelled out. Level 1-1 has nine tiles, including capital.",
                )
            else:
                set_paragraph_text(
                    paragraph,
                    "The Metrics column lists every proving tile on the six Level 1 pages (credit union labels), with abbreviations spelled out. Level 1-1 has nine tiles, including net worth.",
                )
            continue
        new = transform_text(raw, charter)
        if new != raw:
            set_paragraph_text(paragraph, new)

    for paragraph in to_delete:
        delete_paragraph(paragraph)

    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
        is_targets = header[:3] == ["Home report", "Proving metric (bank)", "Sample target"]
        is_metrics_def = header and header[0].startswith("Metric")
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                raw = cell.text
                if is_targets and ri > 0 and charter == "cu":
                    continue
                new = transform_text(raw, charter)
                if is_metrics_def and ri > 0:
                    new = apply_slash_labels(new, charter)
                    new = pick_bank_cu_segments(new, charter)
                if new != raw:
                    cell.text = new
        if is_targets and charter == "cu":
            for ri in range(1, len(table.rows)):
                home, metric, target = CU_TARGETS[ri - 1]
                table.rows[ri].cells[0].text = home
                table.rows[ri].cells[1].text = metric
                table.rows[ri].cells[2].text = target

    _patch_camels_table(doc, charter)
    _patch_nine_tiles_paragraph(doc, charter)
    _patch_summary_table_metrics(doc, charter)
    _patch_glossary_edit_row(doc, charter)
    _scrub_peer_book_mentions(doc, charter)
    _trim_glossary(doc, charter)
    _set_titles(doc, charter)


def _patch_camels_table(doc: Document, charter: str) -> None:
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Letter":
            continue
        for ri in range(1, len(table.rows)):
            cell = table.rows[ri].cells[2]
            t = cell.text
            if charter == "bank":
                t = t.replace("Tier 1 leverage (bank); net worth ratio (credit union).", "Tier 1 leverage.")
                t = t.replace("Credit-union board education hours.", "")
                t = re.sub(r"Return on Equity / net worth growth", "Return on Equity", t)
                t = t.replace("Deposit / share growth", "Deposit growth")
            else:
                t = t.replace("Tier 1 leverage (bank); net worth ratio (credit union).", "Net worth ratio.")
                t = re.sub(r"Return on Equity / net worth growth", "Net worth growth", t)
                t = t.replace("Deposit / share growth", "Share growth")
                t = t.replace("Credit-union board education hours.", "Board education hours.")
            cell.text = t.strip()


def _patch_nine_tiles_paragraph(doc: Document, charter: str) -> None:
    for paragraph in doc.paragraphs:
        if not paragraph.text.startswith("Nine proving tiles"):
            continue
        if charter == "bank":
            set_paragraph_text(
                paragraph,
                "Nine proving tiles (this year’s board pack): Return on Assets; Efficiency ratio; Tier 1 leverage; Customer Satisfaction Index; Retention; Loan growth; Core funding mix; Past due 30 days or more; Engagement. Each tile is an objective with one number — not a catalog of every metric in the institution.",
            )
        else:
            set_paragraph_text(
                paragraph,
                "Nine proving tiles (this year’s board pack): Return on Assets; Efficiency ratio; Net worth ratio; Net Promoter Score; Retention; Loan growth; Core share mix; Past due 30 days or more; Engagement. Each tile is an objective with one number — not a catalog of every metric in the institution.",
            )


def _patch_summary_table_metrics(doc: Document, charter: str) -> None:
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Report name":
            continue
        for row in table.rows[1:]:
            cell = row.cells[4]
            cell.text = apply_slash_labels(transform_text(cell.text, charter), charter)
        if charter == "cu":
            r1 = table.rows[1].cells[4].text
            r1 = r1.replace("Core funding mix", "Core share mix")
            table.rows[1].cells[4].text = r1


def _patch_glossary_edit_row(doc: Document, charter: str) -> None:
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Term":
            continue
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == "Edit (metric objective)":
                if charter == "bank":
                    row.cells[1].text = (
                        "Top-bar pill that changes name, objective, owner, target, peer, unit, or direction "
                        "on the bank charter side of a tile. Session only."
                    )
                else:
                    row.cells[1].text = (
                        "Top-bar pill that changes name, objective, owner, target, peer, unit, or direction "
                        "on the credit union charter side of a tile. Session only."
                    )


def _scrub_peer_book_mentions(doc: Document, charter: str) -> None:
    def scrub(text: str) -> str:
        if charter == "bank":
            return text.replace(
                "Do not mix Financial Performance Report curves on the same tile.",
                "Peer curves use Uniform Bank Performance Report only.",
            )
        return text.replace(
            "Do not mix Uniform Bank Performance Report curves on the same tile.",
            "Peer curves use Financial Performance Report only.",
        ).replace("Uniform Bank Performance Report", "the bank peer book")

    for paragraph in doc.paragraphs:
        new = scrub(paragraph.text)
        if new != paragraph.text:
            set_paragraph_text(paragraph, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                new = scrub(cell.text)
                if new != cell.text:
                    cell.text = new


def _trim_glossary(doc: Document, charter: str) -> None:
    if not doc.tables:
        return
    glossary = doc.tables[-1]
    if glossary.rows[0].cells[0].text.strip() != "Term":
        return
    remove_terms = (
        {"Financial Performance Report", "5300 Call Report", "Field of membership", "Member business lending", "Average Member Relationship", "Net Promoter Score"}
        if charter == "bank"
        else {"Uniform Bank Performance Report", "Call Report", "Customer Satisfaction Index", "Prompt Corrective Action / Community Bank Leverage Ratio"}
    )
    for ri in range(len(glossary.rows) - 1, 0, -1):
        term = glossary.rows[ri].cells[0].text.strip()
        if term in remove_terms:
            glossary._tbl.remove(glossary.rows[ri]._tr)


def _set_titles(doc: Document, charter: str) -> None:
    if not doc.paragraphs:
        return
    if charter == "bank":
        set_paragraph_text(doc.paragraphs[0], "LEVEL 1 EXECUTIVE DASHBOARD SUITE — Bank demo")
        if len(doc.paragraphs) > 1:
            set_paragraph_text(doc.paragraphs[1], "Demo Notes — Community bank")
    else:
        set_paragraph_text(doc.paragraphs[0], "LEVEL 1 EXECUTIVE DASHBOARD SUITE — Credit union demo")
        if len(doc.paragraphs) > 1:
            set_paragraph_text(doc.paragraphs[1], "Demo Notes — Credit union")
        if len(doc.paragraphs) > 2:
            set_paragraph_text(
                doc.paragraphs[2],
                "First walkthrough brief for stakeholders who have not yet seen the Level 1 mockups. Sample institution: Riverbend Community Credit Union (demo label set). As-of date in the mockups: 31 July 2026.",
            )


def build(charter: str, dest: Path) -> None:
    shutil.copy2(MASTER, dest)
    doc = Document(dest)
    transform_document(doc, charter)
    doc.save(dest)


def main() -> None:
    if not MASTER.exists():
        raise SystemExit(f"Missing master copy: {MASTER}")
    build("bank", OUT_BANK)
    build("cu", OUT_CU)
    print(f"Wrote {OUT_BANK.name}")
    print(f"Wrote {OUT_CU.name}")


if __name__ == "__main__":
    main()
